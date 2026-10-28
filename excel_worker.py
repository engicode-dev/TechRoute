import openpyxl
import os


def fill_gost_template(template_path, output_path, header_data, operations_list):
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Шаблон {template_path} не найден!")

    wb = openpyxl.load_workbook(template_path)
    ws1 = wb['МК']

    def safe_write(sheet, cell_ref, value):
        for merged_range in list(sheet.merged_cells.ranges):
            if cell_ref in merged_range:
                min_col, min_row, max_col, max_row = merged_range.bounds
                sheet.unmerge_cells(str(merged_range))
                sheet.cell(row=min_row, column=min_col).value = value
                sheet.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
                return
        sheet[cell_ref] = value

    # 1. ЗАПОЛНЕНИЕ ШАПКИ
    safe_write(ws1, 'J8', "Тюстин С.А.")
    safe_write(ws1, 'AL8', "Учебный кейс")
    safe_write(ws1, 'BB8', header_data['name'])
    safe_write(ws1, 'CO8', header_data['code'])
    safe_write(ws1, 'CY11', "1")
    safe_write(ws1, 'DC11', "1")
    safe_write(ws1, 'F13', header_data['material'])

    # 2. ЗАПОЛНЕНИЕ ОПЕРАЦИЙ
    start_row = 19  # Строка, с которой начинаются операции в твоем шаблоне
    op_counter = 5  # 005, 010, 015...
    row_number = 3  # Порядковый номер строки в таблице (03, 04, 05...)

    for op_data in operations_list:
        # Если не влезаем на страницу — прерываем (в будущем можно добавить перенос на 2-й лист)
        if start_row > 31:
            break

        formatted_op_num = f"{op_counter:03d}"

        # Данные для заполнения (в будущем можно брать из базы, пока заглушки)
        profession = "Слесарь"
        rank = "3"
        description = "1. Установить деталь. 2. Обработать согласно чертежу."
        tool = "Штангенциркуль, ключ гаечный"

        # --- 1. СТРОКА А ---
        safe_write(ws1, f'A{start_row}', 'А')
        safe_write(ws1, f'B{start_row}', f"{row_number:02d}")  # Порядковый номер строки (03, 04...)

        # В зависимости от макета, номер операции может быть в столбце F или V.
        # Судя по эталону, он идет сразу после номера строки.
        safe_write(ws1, f'F{start_row}', formatted_op_num)
        safe_write(ws1, f'V{start_row}', formatted_op_num)

        # Название операции
        op_code = op_data.get("code", "")
        op_full_name = f"{op_code} {op_data['name']}".strip()
        safe_write(ws1, f'AY{start_row}', op_full_name)

        safe_write(ws1, f'BJ{start_row}', profession)
        safe_write(ws1, f'BN{start_row}', rank)

        start_row += 1
        row_number += 1

        # --- 2. СТРОКА Б ---
        if start_row <= 31:
            safe_write(ws1, f'A{start_row}', 'Б')
            safe_write(ws1, f'B{start_row}', f"{row_number:02d}")
            # В эталоне чертеж/эскиз пишется в столбец V. Ставим прочерк.
            safe_write(ws1, f'V{start_row}', "-")
            safe_write(ws1, f'AY{start_row}', op_data["equipment"])

            start_row += 1
            row_number += 1

        # --- 3. СТРОКА О ---
        if start_row <= 31:
            safe_write(ws1, f'A{start_row}', 'О')
            safe_write(ws1, f'B{start_row}', f"{row_number:02d}")
            safe_write(ws1, f'AY{start_row}', description)

            start_row += 1
            row_number += 1

        # --- 4. СТРОКА Т ---
        if start_row <= 31:
            safe_write(ws1, f'A{start_row}', 'Т')
            safe_write(ws1, f'B{start_row}', f"{row_number:02d}")
            safe_write(ws1, f'AY{start_row}', tool)

            start_row += 1
            row_number += 1

        # --- 5. ПУСТАЯ СТРОКА (РАЗДЕЛИТЕЛЬ) ---
        if start_row <= 31:
            # Оставляем строку пустой, просто переходим дальше
            start_row += 1
            row_number += 1

        op_counter += 5

    wb.save(output_path)
    return True