import openpyxl
import os


def write_to_excel(template_path, output_path, header_data, prepared_operations):
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Шаблон не найден: {template_path}")

    wb = openpyxl.load_workbook(template_path)
    ws1 = wb['МК']
    ws2 = wb['МКа']

    def safe_write(sheet, cell_ref, value):
        for merged_range in list(sheet.merged_cells.ranges):
            if cell_ref in merged_range:
                min_col, min_row, max_col, max_row = merged_range.bounds
                sheet.unmerge_cells(str(merged_range))
                sheet.cell(row=min_row, column=min_col).value = value
                sheet.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
                return
        sheet[cell_ref] = value

    # --- ШАПКА ---
    safe_write(ws1, 'J8', header_data.get('developer', 'Инженер')) # Разраб
    safe_write(ws1, 'AL8', header_data.get('organization', 'Организация')) # Организация
    safe_write(ws1, 'BB8', header_data['name'])
    safe_write(ws1, 'CO8', header_data['code'])
    safe_write(ws1, 'F13', header_data['material'])

    # Записываем массу заготовки в графу МЗ
    if 'mass' in header_data:
        safe_write(ws1, 'CU15', str(header_data['mass']))

    row_page1 = 19
    row_page2 = 14
    op_counter = 5
    row_num = 3

    page_turned = False
    total_pages = 1

    def write_row(prefix, op_data=None):
        nonlocal row_page1, row_page2, row_num, page_turned, total_pages

        if not page_turned and row_page1 > 31:
            page_turned = True
            total_pages = 2

        sheet = ws1 if not page_turned else ws2
        current_row = row_page1 if not page_turned else row_page2

        safe_write(sheet, f'A{current_row}', f"{prefix} {row_num:02d}")

        if prefix == "А":
            safe_write(sheet, f'F{current_row}', "1")
            safe_write(sheet, f'J{current_row}', "1")
            safe_write(sheet, f'R{current_row}', "1")
            safe_write(sheet, f'V{current_row}', f"{op_counter:03d}")
            safe_write(sheet, f'AY{current_row}', f"{op_data['code']} {op_data['name']}".strip())
            safe_write(sheet, f'BJ{current_row}', op_data['profession'])
            safe_write(sheet, f'BN{current_row}', op_data['rank'])
            safe_write(sheet, f'CR{current_row}', op_data['time_Tsht'])
        elif prefix == "Б":
            safe_write(sheet, f'AY{current_row}', op_data['equipment'])
        elif prefix == "О":
            safe_write(sheet, f'AY{current_row}', op_data['description'])
        elif prefix == "Т":
            safe_write(sheet, f'AY{current_row}', op_data['tool'])

        if not page_turned:
            row_page1 += 1
        else:
            row_page2 += 1
        row_num += 1

    for op in prepared_operations:
        write_row("А", op)
        write_row("Б", op)
        write_row("О", op)
        write_row("Т", op)
        write_row("-", None)

        op_counter += 5

    safe_write(ws1, 'CY11', str(total_pages))
    if page_turned:
        safe_write(ws2, 'BS8', header_data['code'])
        safe_write(ws2, 'DA6', "2")

    # --- СКРЫВАЕМ ЛИШНИЕ ЛИСТЫ, ЧТОБЫ ОНИ НЕ ПОПАДАЛИ В PDF ---
    for sheet in wb.sheetnames:
        if sheet not in ['МК', 'МКа']:
            wb[sheet].sheet_state = 'hidden'

    # Скрываем вторую страницу, если мы не переходили на неё
    if not page_turned:
        wb['МКа'].sheet_state = 'hidden'
    else:
        wb['МКа'].sheet_state = 'visible'

    wb.save(output_path)