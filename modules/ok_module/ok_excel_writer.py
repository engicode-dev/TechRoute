import openpyxl
import os


def write_to_ok(template_path, output_path, header_data, transitions_list, total_main_time, aux_time, final_time):
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Шаблон ОК не найден: {template_path}")

    wb = openpyxl.load_workbook(template_path)
    ws1 = wb['ОК']
    ws2 = wb['ОКа']

    def safe_write(sheet, cell_ref, value):
        for merged_range in list(sheet.merged_cells.ranges):
            if cell_ref in merged_range:
                min_col, min_row, max_col, max_row = merged_range.bounds
                sheet.unmerge_cells(str(merged_range))
                sheet.cell(row=min_row, column=min_col).value = value
                sheet.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
                return
        sheet[cell_ref] = value

    # --- ШАПКА (Первая страница) ---
    safe_write(ws1, 'J8', header_data.get('developer', 'Инженер'))
    safe_write(ws1, 'AL8', header_data.get('organization', 'Организация'))
    safe_write(ws1, 'AL11', header_data['name'])
    safe_write(ws1, 'BB8', header_data['code'])
    safe_write(ws1, 'CU6', "1")  # Всего страниц (пока 1, потом поменяем)
    safe_write(ws1, 'DA6', "1")  # Текущая страница
    safe_write(ws1, 'AE14', header_data['material'])
    safe_write(ws1, 'BD14', str(header_data.get('hardness_hb', '')))
    safe_write(ws1, 'A14', header_data.get('op_name', ''))
    safe_write(ws1, 'A16', header_data.get('op_equipment', ''))

    # Время в шапку
    safe_write(ws1, 'BD16', f"{total_main_time:.2f}")
    safe_write(ws1, 'BK16', f"{aux_time:.2f}")
    safe_write(ws1, 'BZ16', f"{final_time:.2f}")

    # --- ТАБЛИЦА ПЕРЕХОДОВ ---
    start_row = 18
    max_row = 30
    trans_num = 1
    page_turned = False

    for trans in transitions_list:
        # Если первая страница заполнена, переключаемся на вторую
        if start_row > max_row and not page_turned:
            page_turned = True

        current_sheet = ws1 if not page_turned else ws2
        current_row = start_row if not page_turned else (start_row - max_row + 11)  # Сдвиг для второй страницы

        if page_turned and current_row > 17:  # Ограничение второй страницы (по твоему макету до 17)
            break  # Если и на второй не влезло — обрезаем (в будущем можно добавить третью)

        # 1. Строка О
        safe_write(current_sheet, f'A{current_row}', f"О {trans_num:02d}")
        safe_write(current_sheet, f'F{current_row}', trans.get('text', ''))
        current_row += 1

        # 2. Строка Т
        if trans.get('tool'):
            safe_write(current_sheet, f'A{current_row}', 'Т')
            safe_write(current_sheet, f'F{current_row}', trans.get('tool', ''))
            current_row += 1

        # 3. Строка Р (Режимы)
        params = trans.get('params', {})
        if any(params.values()):
            safe_write(current_sheet, f'A{current_row}', 'Р')
            if params.get('D'): safe_write(current_sheet, f'BC{current_row}', params['D'])
            if params.get('L'): safe_write(current_sheet, f'BN{current_row}', params['L'])
            if params.get('t'): safe_write(current_sheet, f'BW{current_row}', params['t'])
            if params.get('i'): safe_write(current_sheet, f'CB{current_row}', params['i'])
            if params.get('S'): safe_write(current_sheet, f'CH{current_row}', params['S'])
            if params.get('n'): safe_write(current_sheet, f'CR{current_row}', params['n'])
            if params.get('V'): safe_write(current_sheet, f'CY{current_row}', params['V'])
            current_row += 1

        # Пустая строка
        if trans_num < len(transitions_list):
            current_row += 1

        start_row = current_row
        trans_num += 1

    # Обновляем шапку (если была вторая страница)
    if page_turned:
        safe_write(ws1, 'CU6', "2")  # Всего страниц
        safe_write(ws2, 'BN8', header_data['code'])
        safe_write(ws2, 'DA6', "2")  # Номер текущей

    # --- СКРЫВАЕМ ЛИШНИЕ ЛИСТЫ ДЛЯ PDF ---
    for sheet in wb.sheetnames:
        if sheet not in ['ОК', 'ОКа']:
            wb[sheet].sheet_state = 'hidden'

    if not page_turned:
        wb['ОКа'].sheet_state = 'hidden'

    wb.save(output_path)
    return True