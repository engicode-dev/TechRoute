import openpyxl
import os
from openpyxl.styles import Font


def write_spec_excel(template_path, output_path, header_data, items_list):
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Шаблон спецификации не найден: {template_path}")

    # Загружаем шаблон
    wb = openpyxl.load_workbook(template_path)
    ws1 = wb['SP']

    def safe_write(sheet, cell_ref, value):
        for merged_range in list(sheet.merged_cells.ranges):
            if cell_ref in merged_range:
                min_col, min_row, max_col, max_row = merged_range.bounds
                sheet.unmerge_cells(str(merged_range))
                sheet.cell(row=min_row, column=min_col).value = value
                sheet.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)
                return
        sheet[cell_ref] = value

    # --- ШАПКА (Штамп) ---
    safe_write(ws1, 'R263', header_data.get('developer', 'Инженер'))
    safe_write(ws1, 'BN248', header_data['code'])
    safe_write(ws1, 'BN263', header_data['name'])

    safe_write(ws1, 'EF268', "У")
    safe_write(ws1, 'EK268', "")
    safe_write(ws1, 'EP268', "")
    safe_write(ws1, 'EU268', "1")
    safe_write(ws1, 'FJ268', "1")

    safe_write(ws1, 'EF273', header_data.get('organization', 'Организация'))

    # --- ТАБЛИЦА ---
    start_row = 16
    step = 8
    max_row = 240

    categories_order = ["Документация", "Сборочные единицы", "Детали", "Стандартные изделия", "Материалы"]
    grouped_items = {cat: [] for cat in categories_order}

    for item in items_list:
        cat = item.get('category', 'Детали')
        if cat in grouped_items:
            grouped_items[cat].append(item)
        else:
            grouped_items['Детали'].append(item)

    current_row = start_row
    pos_counter = 1

    for category in categories_order:
        items = grouped_items[category]
        if not items:
            continue

        # Пустая строка перед разделом
        current_row += step
        if current_row > max_row:
            break

        # Заголовок раздела
        safe_write(ws1, f'CM{current_row}', category)
        cell = ws1[f'CM{current_row}']
        # Шрифт с подчеркиванием
        cell.font = Font(name='GOSTRUS', size=10, bold=True, underline='single')

        current_row += step
        current_row += step  # Пустая строка после раздела

        for item in items:
            if current_row > max_row:
                break

            if category != "Документация":
                safe_write(ws1, f'M{current_row}', str(pos_counter))
                pos_counter += 1

            safe_write(ws1, f'A{current_row}', item.get('format', 'А4'))
            safe_write(ws1, f'G{current_row}', item.get('zone', ''))
            safe_write(ws1, f'U{current_row}', item.get('code', ''))
            safe_write(ws1, f'CM{current_row}', item.get('name', ''))
            safe_write(ws1, f'EX{current_row}', str(item.get('count', 1)))

            current_row += step

    wb.save(output_path)
    return True